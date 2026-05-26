import asyncio
from typing import Dict, List, Any, Optional, AsyncGenerator
from dataclasses import dataclass, field
from enum import Enum
import json
import os
from pathlib import Path

from backend.core.llm_manager import llm_manager
from backend.core.tool_watcher import tool_registry

class NodeType(Enum):
    AGENT = "agent"
    TOOL = "tool"
    JUDGE = "judge"
    LOOP = "loop"
    CONDITION = "condition"
    INPUT = "input"
    OUTPUT = "output"
    RAG = "rag"

@dataclass
class Node:
    id: str
    type: NodeType
    config: Dict[str, Any] = field(default_factory=dict)
    position: Dict[str, float] = field(default_factory=dict)

@dataclass
class Edge:
    source: str
    target: str
    condition: Optional[str] = None
    label: str = ""

@dataclass
class Pipeline:
    id: str
    name: str
    nodes: List[Node]
    edges: List[Edge]
    variables: Dict[str, Any] = field(default_factory=dict)

class PipelineEngine:
    def __init__(self):
        self.pipelines: Dict[str, Pipeline] = {}
        self._running: Dict[str, bool] = {}

    def create_pipeline(self, pipeline_data: dict) -> Pipeline:
        nodes = [Node(
            id=n["id"],
            type=NodeType(n["type"]),
            config=n.get("config", {}),
            position=n.get("position", {})
        ) for n in pipeline_data.get("nodes", [])]

        edges = [Edge(
            source=e["source"],
            target=e["target"],
            condition=e.get("condition"),
            label=e.get("label", "")
        ) for e in pipeline_data.get("edges", [])]

        pipeline = Pipeline(
            id=pipeline_data["id"],
            name=pipeline_data.get("name", "Untitled"),
            nodes=nodes,
            edges=edges
        )
        self.pipelines[pipeline.id] = pipeline
        return pipeline

    async def execute(self, pipeline_id: str, initial_input: str = "", context: Dict = None) -> AsyncGenerator[Dict, None]:
        pipeline = self.pipelines.get(pipeline_id)
        if not pipeline:
            yield {"error": "Pipeline not found"}
            return

        self._running[pipeline_id] = True
        state = {"input": initial_input, "output": "", "variables": context or {}, "history": []}

        graph: Dict[str, List[Edge]] = {}
        for edge in pipeline.edges:
            if edge.source not in graph:
                graph[edge.source] = []
            graph[edge.source].append(edge)

        start_nodes = [n for n in pipeline.nodes if n.type == NodeType.INPUT]
        if not start_nodes:
            start_nodes = pipeline.nodes[:1]

        current_nodes = start_nodes
        visited = set()
        max_steps = 50
        step = 0

        while current_nodes and step < max_steps and self._running.get(pipeline_id, False):
            step += 1
            next_nodes = []

            for node in current_nodes:
                if node.id in visited:
                    continue
                visited.add(node.id)

                yield {"step": step, "node": node.id, "type": node.type.value, "status": "running"}

                result = await self._execute_node(node, state, pipeline)
                state["output"] = result
                state["history"].append({"node": node.id, "output": result})

                yield {"step": step, "node": node.id, "type": node.type.value, "status": "completed", "output": result}

                outgoing = graph.get(node.id, [])
                for edge in outgoing:
                    if edge.condition:
                        if self._eval_condition(edge.condition, state):
                            target = next((n for n in pipeline.nodes if n.id == edge.target), None)
                            if target:
                                next_nodes.append(target)
                    else:
                        target = next((n for n in pipeline.nodes if n.id == edge.target), None)
                        if target:
                            next_nodes.append(target)

            current_nodes = next_nodes

        self._running[pipeline_id] = False
        yield {"status": "finished", "final_output": state["output"], "history": state["history"]}

    async def _execute_node(self, node: Node, state: Dict, pipeline: Pipeline) -> str:
        config = node.config

        if node.type == NodeType.INPUT:
            source_type = config.get("source_type", "text")
            if source_type == "file":
                file_path = config.get("file_path", "")
                file_content = config.get("file_content", "")
                if file_content:
                    state["input"] = file_content
                    return f"[File: {config.get('file_name', 'unknown')}]\n{file_content[:500]}..."
                elif file_path:
                    try:
                        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                            content = f.read()
                        state["input"] = content
                        return f"[File: {file_path}]\n{content[:500]}..."
                    except Exception as e:
                        return f"Error reading file: {e}"
            elif source_type == "url":
                url = config.get("file_path", "")
                try:
                    import httpx
                    async with httpx.AsyncClient(timeout=10.0) as client:
                        resp = await client.get(url)
                        content = resp.text[:5000]
                    state["input"] = content
                    return f"[URL: {url}]\n{content[:500]}..."
                except Exception as e:
                    return f"Error fetching URL: {e}"
            else:
                default_val = config.get("default_value", "")
                state["input"] = default_val
                return default_val

        elif node.type == NodeType.AGENT:
            system_prompt = config.get("system_prompt", "Ty poleznyy assistent.")
            temperature = config.get("temperature", 0.3)
            max_tokens = config.get("max_tokens", 4096)
            timeout = config.get("timeout", 60)
            retry_count = config.get("retry_count", 1)

            messages = [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": state["input"]}
            ]

            if config.get("use_rag"):
                from backend.rag.vector_store import vector_store
                if vector_store._model_loaded:
                    rag_results = await vector_store.similarity_search(state["input"], k=config.get("rag_top_k", 3))
                    context_text = "\n\n".join([r["content"] for r in rag_results])
                    messages.insert(1, {"role": "system", "content": f"Kontekst iz dokumentov:\n{context_text}"})

            tools = None
            if config.get("tools"):
                tools = [tool_registry.get(t) for t in config["tools"] if tool_registry.get(t)]

            result = ""
            for attempt in range(retry_count + 1):
                try:
                    async for chunk in llm_manager.chat(messages, tools=tools, stream=False):
                        result += chunk
                    break
                except Exception as e:
                    if attempt >= retry_count:
                        return f"Error after {retry_count + 1} attempts: {e}"
                    await asyncio.sleep(1)
            return result

        elif node.type == NodeType.TOOL:
            tool_name = config.get("tool_name", "")
            tool_func = tool_registry.get(tool_name)
            if tool_func:
                params = config.get("parameters", {})
                for key, val in params.items():
                    if isinstance(val, str) and val.startswith("$"):
                        params[key] = state["variables"].get(val[1:], val)

                timeout = config.get("timeout", 30)
                retry_on_error = config.get("retry_on_error", True)
                fallback = config.get("fallback_value", "")

                try:
                    if asyncio.iscoroutinefunction(tool_func):
                        result = await asyncio.wait_for(tool_func(**params), timeout=timeout)
                    else:
                        result = tool_func(**params)
                    return str(result)
                except Exception as e:
                    if retry_on_error:
                        try:
                            await asyncio.sleep(0.5)
                            if asyncio.iscoroutinefunction(tool_func):
                                result = await asyncio.wait_for(tool_func(**params), timeout=timeout)
                            else:
                                result = tool_func(**params)
                            return str(result)
                        except:
                            return fallback or f"Tool error: {e}"
                    return fallback or f"Tool error: {e}"
            return f"Tool '{tool_name}' not found"

        elif node.type == NodeType.JUDGE:
            criteria = config.get("criteria", "Kachestvo otveta")
            min_score = config.get("min_score", 7)
            feedback_mode = config.get("feedback_mode", "score_only")

            result = await llm_manager.judge(
                criteria=criteria,
                input_text=state["input"],
                output_text=state["output"],
                context=json.dumps(state["history"], ensure_ascii=False)
            )

            result["min_score"] = min_score
            result["passed"] = result.get("score", 0) >= min_score

            if feedback_mode == "score_only":
                return str(result.get("score", 0))
            elif feedback_mode == "score_and_reason":
                return f"Оценка: {result.get('score', 0)}\nПричина: {result.get('reason', '')}"
            return json.dumps(result, ensure_ascii=False)

        elif node.type == NodeType.LOOP:
            iterations = config.get("iterations", 3)
            break_on_empty = config.get("break_on_empty", True)
            break_on_error = config.get("break_on_error", False)
            collect_results = config.get("collect_results", True)
            delay_ms = config.get("delay_ms", 0)

            results = []
            for i in range(iterations):
                if not self._running.get(pipeline.id, False):
                    break
                state["variables"]["iteration"] = i
                try:
                    if delay_ms > 0:
                        await asyncio.sleep(delay_ms / 1000)
                    if collect_results:
                        results.append(f"[{i}] {state['output'][:200]}")
                    else:
                        state["output"] = f"Iteration {i}: {state['output']}"
                except Exception as e:
                    if break_on_error:
                        results.append(f"[{i}] ERROR: {e}")
                        break
                    results.append(f"[{i}] ERROR: {e}")

                if break_on_empty and not state["output"].strip():
                    break

            return "\n".join(results) if collect_results else state["output"]

        elif node.type == NodeType.CONDITION:
            condition = config.get("condition", "true")
            case_sensitive = config.get("case_sensitive", False)

            try:
                safe_dict = {"output": state["output"], "input": state["input"], **state["variables"]}
                if not case_sensitive and isinstance(state["output"], str):
                    safe_dict["output_lower"] = state["output"].lower()
                result = eval(condition, {"__builtins__": {}}, safe_dict)
                return str(result)
            except:
                return "true"

        elif node.type == NodeType.RAG:
            from backend.rag.vector_store import vector_store
            if not vector_store._model_loaded:
                return "RAG model not loaded"

            query = state["input"]
            k = config.get("top_k", 5)
            score_threshold = config.get("score_threshold", 0.0)
            include_metadata = config.get("include_metadata", True)

            results = await vector_store.similarity_search(query, k=k)
            if score_threshold > 0:
                results = [r for r in results if r.get("score", 0) >= score_threshold]

            if not include_metadata:
                for r in results:
                    r.pop("metadata", None)

            return json.dumps(results, ensure_ascii=False)

        elif node.type == NodeType.OUTPUT:
            template = config.get("template", "{{output}}")
            output_text = template.replace("{{output}}", state["output"])
            output_text = output_text.replace("{{input}}", state["input"])

            if config.get("show_timestamp"):
                from datetime import datetime
                output_text = f"[{datetime.now().isoformat()}]\n{output_text}"

            preview_len = config.get("preview_length", 500)
            preview = output_text[:preview_len]

            if config.get("save_to_file"):
                file_path = config.get("file_path", "")
                file_format = config.get("file_format", "txt")
                append_mode = config.get("append_mode", False)

                if file_path:
                    try:
                        p = Path(file_path)
                        p.parent.mkdir(parents=True, exist_ok=True)

                        if file_format == "csv":
                            import csv
                            mode = "a" if append_mode else "w"
                            with open(p, mode, newline="", encoding="utf-8") as f:
                                writer = csv.writer(f)
                                lines = output_text.strip().split("\n")
                                for line in lines:
                                    writer.writerow([line])
                        elif file_format == "xlsx":
                            try:
                                from openpyxl import Workbook, load_workbook
                                if append_mode and p.exists():
                                    wb = load_workbook(p)
                                    ws = wb.active
                                    next_row = ws.max_row + 1
                                else:
                                    wb = Workbook()
                                    ws = wb.active
                                    next_row = 1
                                ws.cell(row=next_row, column=1, value=output_text)
                                wb.save(p)
                            except ImportError:
                                return "Error: openpyxl not installed. pip install openpyxl"
                        elif file_format == "json":
                            mode = "a" if append_mode else "w"
                            with open(p, mode, encoding="utf-8") as f:
                                json.dump({"output": output_text, "input": state["input"]}, f, ensure_ascii=False)
                                f.write("\n")
                        else:
                            mode = "a" if append_mode else "w"
                            with open(p, mode, encoding="utf-8") as f:
                                f.write(output_text + "\n")

                        return f"[Saved to {file_path}]\n{preview}"
                    except Exception as e:
                        return f"[Save error: {e}]\n{preview}"

            return preview

        return state["input"]

    def _eval_condition(self, condition: str, state: Dict) -> bool:
        try:
            safe_dict = {"output": state["output"], "input": state["input"], **state["variables"]}
            return eval(condition, {"__builtins__": {}}, safe_dict)
        except:
            return True

    def stop_pipeline(self, pipeline_id: str):
        self._running[pipeline_id] = False
